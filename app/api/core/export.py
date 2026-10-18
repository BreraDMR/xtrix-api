import csv
import io
from typing import List, Dict, Any, Optional
from datetime import datetime
from fastapi import Response
try:
    import openpyxl  # noqa: F401 - just checking it is installed
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

_FORMULA_START = ("=", "+", "-", "@")


def _csv_text(value: Any) -> Any:
    # Excel/Sheets run a cell starting with = + - @ (or tab/CR) as a formula, and they skip
    # leading whitespace before looking. A leading ' makes it plain text. Numbers stay numbers.
    if not isinstance(value, str):
        return value
    if value.startswith(("\t", "\r")) or value.lstrip().startswith(_FORMULA_START):
        return "'" + value
    return value


def _excel_text_row(ws, values):
    ws.append(values)
    # User text beginning with '=' must stay text, including column headings.
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def export_to_csv(
    data: List[Dict[str, Any]],
    filename: str = "export.csv",
    fieldnames: Optional[List[str]] = None,
) -> Response:
    output = io.StringIO()
    if not data:
        writer = csv.DictWriter(output, fieldnames=fieldnames or [])
        writer.writerow({key: _csv_text(key) for key in writer.fieldnames})
    else:
        if fieldnames is None:
            fieldnames = list(data[0].keys())
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writerow({key: _csv_text(key) for key in writer.fieldnames})
        for row in data:
            flat_row = {}
            for key in fieldnames:
                value = row.get(key)
                if isinstance(value, dict):
                    flat_row[key] = str(value)
                elif isinstance(value, datetime):
                    flat_row[key] = value.isoformat()
                elif value is None:
                    flat_row[key] = ""
                else:
                    flat_row[key] = value
            writer.writerow({key: _csv_text(value) for key, value in flat_row.items()})
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
def export_to_excel(
    data: List[Dict[str, Any]],
    filename: str = "export.xlsx",
    sheet_name: str = "Data",
    fieldnames: Optional[List[str]] = None,
) -> Response:
    if not EXCEL_AVAILABLE:
        raise ImportError(
            "openpyxl is required for Excel export. Install it with: pip install openpyxl"
        )
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    if not data:
        if fieldnames:
            _excel_text_row(ws, fieldnames)
    else:
        if fieldnames is None:
            fieldnames = list(data[0].keys())
        _excel_text_row(ws, fieldnames)
        # ws.append returns None, style the row it just wrote instead
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in data:
            row_data = []
            for key in fieldnames:
                value = row.get(key)
                if isinstance(value, dict):
                    row_data.append(str(value))
                elif isinstance(value, datetime):
                    row_data.append(value.isoformat())
                elif value is None:
                    row_data.append("")
                else:
                    row_data.append(value)
            _excel_text_row(ws, row_data)
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
def _as_dict(obj: Any) -> Dict[str, Any]:
    # endpoints hand us ORM rows, pydantic models or plain dicts - dict(orm_row) used to blow up
    if isinstance(obj, dict):
        return obj
    if hasattr(obj, 'model_dump'):
        return obj.model_dump()
    if hasattr(obj, '__table__'):
        return {c.name: getattr(obj, c.name) for c in obj.__table__.columns}
    return dict(obj)
def prepare_client_export_data(clients: List[Any]) -> List[Dict[str, Any]]:
    export_data = []
    for client in clients:
        data = _as_dict(client)
        export_data.append({
            'ID': data.get('id'),
            'Full Name': data.get('full_name'),
            'Birth Date': data.get('birth_date'),
            'Gender': data.get('gender'),
            'Contact Phone': data.get('contact_phone'),
            'Contact Email': data.get('contact_email'),
            'Allergies': data.get('allergies'),
            'Chronic Diseases': data.get('chronic_diseases'),
            'Pregnancy': data.get('pregnancy'),
            'Viral Diseases': data.get('viral_diseases'),
            'Spending': data.get('spending'),
            'Birthday Discount': data.get('birthday_discount'),
            'Created At': data.get('created_at'),
        })
    return export_data
def prepare_visit_export_data(visits: List[Any]) -> List[Dict[str, Any]]:
    export_data = []
    for visit in visits:
        data = _as_dict(visit)
        export_data.append({
            'ID': data.get('id'),
            'Client ID': data.get('client_id'),
            'Specialist ID': data.get('specialist_id'),
            'Service ID': data.get('service_id'),
            'Date Start': data.get('date_start'),
            'Date End': data.get('date_end'),
            'Status': data.get('status'),
            'Price': data.get('price'),
            'Notes': data.get('note') or data.get('notes'),
            'Created At': data.get('created_at'),
        })
    return export_data
def prepare_service_export_data(services: List[Any]) -> List[Dict[str, Any]]:
    export_data = []
    for service in services:
        data = _as_dict(service)
        export_data.append({
            'ID': data.get('id'),
            'Specialist ID': data.get('specialist_id'),
            'Name': data.get('name'),
            'Description': data.get('description'),
            'Duration Minutes': data.get('duration_minutes'),
            'Price': data.get('price'),
            'Is Active': data.get('is_active'),
            'Created At': data.get('created_at'),
        })
    return export_data
