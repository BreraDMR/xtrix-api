import csv
import io

import pytest
from openpyxl import load_workbook

from app.api.core.export import export_to_csv, export_to_excel
from tests.tenant_world import build_tenant


@pytest.mark.parametrize('value', ['=1+1', '+420123456789', '-1+2', '@SUM(A1)', '\t=1+1', ' \r=1+1', '\tplain', '\rplain'])
def test_csv_cells_cannot_be_formulas(value):
    response = export_to_csv([{'Name': value, 'Amount': -12.5}])
    rows = list(csv.reader(io.StringIO(response.body.decode())))
    assert rows[1] == ["'" + value, '-12.5']


def test_csv_headers_cannot_be_formulas():
    response = export_to_csv([{'=1+1': 'ordinary'}])
    rows = list(csv.reader(io.StringIO(response.body.decode())))
    assert rows == [["'=1+1"], ['ordinary']]


def test_excel_text_is_never_a_formula():
    response = export_to_excel([{'=Heading': '=HYPERLINK("https://example.com")', 'Amount': -5}])
    sheet = load_workbook(io.BytesIO(response.body), data_only=False).active
    assert sheet['A1'].data_type == 's'
    assert sheet['A1'].value == '=Heading'
    assert sheet['A2'].data_type == 's'
    assert sheet['A2'].value.startswith('=HYPERLINK')
    assert sheet['B2'].value == -5
    assert sheet['B2'].data_type == 'n'


def test_plain_text_is_left_alone():
    response = export_to_csv([{'Name': 'Anna = Bob', 'Note': ' spaced'}])
    rows = list(csv.reader(io.StringIO(response.body.decode())))
    assert rows[1] == ['Anna = Bob', ' spaced']


EVIL = '=HYPERLINK("http://evil.example","click")'


def test_client_name_formula_is_neutralised_in_real_exports(client, make_user, headers_for):
    a = build_tenant(client, make_user, headers_for, "Alpha", "+42070011")
    r = client.put(f"/clients/{a['client']['id']}", json={"full_name": EVIL}, headers=a["headers"])
    assert r.status_code == 200, r.text[:300]

    csv_resp = client.get("/clients/export/csv", headers=a["headers"])
    assert csv_resp.status_code == 200
    names = [row[1] for row in csv.reader(io.StringIO(csv_resp.text))]
    assert "'" + EVIL in names
    assert EVIL not in names

    xlsx = client.get("/clients/export/excel", headers=a["headers"])
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    cells = [c for c in sheet['B'] if c.value == EVIL]
    assert cells and all(c.data_type == 's' for c in cells)

    # the analytics report builds its own rows, check it goes through the same escaping
    done = client.put(f"/visits/{a['visit']['id']}", json={"status": "completed", "payment_method": "cash"}, headers=a["headers"])
    assert done.status_code == 200, done.text[:300]
    report = client.post("/analytics/reports/export", json={"report_type": "clients", "format": "csv"}, headers=a["headers"])
    assert report.status_code == 200, report.text[:300]
    report_names = [row[0] for row in csv.reader(io.StringIO(report.text))]
    assert "'" + EVIL in report_names, report_names
