import io

import pytest
from openpyxl import load_workbook

from tests.tenant_world import build_tenant


@pytest.mark.parametrize("kind", ["clients", "services", "visits"])
def test_exports_work_with_real_rows(client, make_user, headers_for, kind):
    a = build_tenant(client, make_user, headers_for, "Alpha", "+42070011")
    csv = client.get(f"/{kind}/export/csv", headers=a["headers"])
    assert csv.status_code == 200, csv.text[:300]
    xlsx = client.get(f"/{kind}/export/excel", headers=a["headers"])
    assert xlsx.status_code == 200, xlsx.text[:300]
    rows = list(load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert len(rows) >= 2 and "Alpha" in str(rows), rows
    assert "Alpha" in csv.text
